# -*- coding: utf-8 -*-
"""scripts/v8_1/12_manuscript_v2.py -- 稿件 v2：出版级叙事重写（EN + ZH）+ 补充材料。

关键变化（按审稿意见）：
* 摘要改为**非结构化**科学叙述（无 Background/Problem/… 标签）
* 引言 4 段；Results 6 节；Discussion 7 段；Limitations 精简为 5 段
* **移除全部 markdown token**（`**` `*not*` 等），改用真实 DOCX 格式
* 移除内部审计语言（GREEN/RED、rowset hash、H1–H5、脚本名、§号）
* 图注改写为论文风格；主体表压缩为 3 张
* 方程用真实上下标与斜体变量
* 新增标题页、作者/单位占位、关键词
* 参考文献：仅保留**已核实**条目；未核实者标 `REFERENCE_VERIFY_REQUIRED`
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # noqa: BLE001
    pass

R = Path(__file__).resolve().parent.parent.parent
A = R / "v8_1"
F = R / "v8_1_final"
V8 = R / "experiments" / "v8"
OUT = R / "v8_1_final" / "manuscript_v2"
OUT.mkdir(parents=True, exist_ok=True)
FM = F / "figures_main"
FS = F / "figures_supplement"

FR = json.load(open(A / "FROZEN_RESULTS.json", encoding="utf-8"))
PU, H4, GEO = FR["primary_universe"], FR["H4_controlled_split"], FR["axis_c_geometry"]
DL = H4.get("delta", {})
ST = pd.read_csv(A / "rank_agreement_by_stratum.csv")
EP = pd.read_csv(A / "external_clean_profiles.csv")
W2 = pd.read_csv(A / "WEIGHTING_EFFECT_SIZE_AUDIT.csv")
UR = pd.read_csv(A / "MATCHED_UNIVERSE_REGISTRY.csv")
er = FR["external_clean"]["rank_agreement_external"]
rmn, rmx = float(ST.rho_overall_vs_macro.min()), float(ST.rho_overall_vs_macro.max())
sf0 = int(W2["signflip_C_B_gt0"].sum()); sf5 = int(W2["signflip_C_B_gt0.005"].sum())
sfw = int(W2["signflip_C_W_gt0"].sum())
best = EP.loc[EP.overall_pearson.idxmax()]
r0, r3 = H4["random"]["overall_pearson"], H4["family_OOD"]["overall_pearson"]
cb0, cw0 = H4["random"]["C_B"], H4["random"]["C_W"]
cb3, cw3 = H4["family_OOD"]["C_B"], H4["family_OOD"]["C_W"]
print(f"r {r0:.4f}->{r3:.4f}; dCB {DL['delta_C_B']:+.4f}; dCW {DL['delta_C_W']:+.4f}; "
      f"between {100*DL['fraction_between']:.1f}%")

# ---------------- references (verified only) ----------------
REFS = [
 ("Jankauskaitė J, Jiménez-García B, Dapkūnas J, Fernández-Recio J, Moal IH. "
  "SKEMPI 2.0: an updated benchmark of changes in protein–protein binding energy, "
  "kinetics and thermodynamics upon mutation. Bioinformatics. 2019;35(3):462–469. "
  "doi:10.1093/bioinformatics/bty635", True),
 ("Bushuiev A, Bushuiev R, Sedlar J, Pluskal T, Damborsky J, Mazurenko S, Sivic J. "
  "Revealing data leakage in protein interaction benchmarks. arXiv:2404.10457. 2024.",
  True),
 ("Bushuiev A, Bushuiev R, Kouba P, Filkin A, Gabrielova M, Gabriel M, Sedlar J, "
  "Pluskal T, Damborsky J, Mazurenko S, Sivic J. Learning to design protein–protein "
  "interactions with enhanced generalization. In: The Twelfth International Conference "
  "on Learning Representations (ICLR). 2024. arXiv:2310.18515", True),
 ("Yu G, Bi X, Ma T, Li Y, Wang J. CATH-ddG: towards robust mutation effect prediction "
  "on protein–protein interactions out of CATH homologous superfamily. Bioinformatics. "
  "2025;41(Suppl_1):i362–i372. doi:10.1093/bioinformatics/btaf228", True),
 ("Barlow KA, Ó Conchúir S, Thompson S, Suresh P, Lucas JE, Heinonen M, Kortemme T. "
  "Flex ddG: Rosetta ensemble-based estimation of changes in protein–protein binding "
  "affinity upon mutation. J Phys Chem B. 2018;122(21):5389–5399. "
  "doi:10.1021/acs.jpcb.7b11367", True),
 ("Delgado J, Radusky LG, Cianferoni D, Serrano L. FoldX 5.0: working with RNA, small "
  "molecules and a new graphical interface. Bioinformatics. 2019;35(20):4168–4169. "
  "doi:10.1093/bioinformatics/btz184", True),
 ("Dauparas J, Anishchenko I, Bennett N, Bai H, Ragotte RJ, Milles LF, et al. Robust "
  "deep learning-based protein sequence design using ProteinMPNN. Science. "
  "2022;378(6615):49–56. doi:10.1126/science.add2187", True),
 ("Hsu C, Verkuil R, Liu J, Lin Z, Hie B, Sercu T, Lerer A, Rives A. Learning inverse "
  "folding from millions of predicted structures. In: Proceedings of the 39th "
  "International Conference on Machine Learning (ICML). 2022;162:8946–8970.", True),
 ("Liu X, Luo Y, Li P, Song S, Peng J. Deep geometric representations for modeling "
  "effects of mutations on protein–protein binding affinity. PLoS Comput Biol. "
  "2021;17(8):e1009284. doi:10.1371/journal.pcbi.1009284", True),
 ("Luo S, Su Y, Wu Z, Su C, Peng J, Ma J. Rotamer density estimator is an unsupervised "
  "learner of the effect of mutations on protein–protein interaction. In: The Eleventh "
  "International Conference on Learning Representations (ICLR). 2023.", True),
 ("Sora V, Laspiur AO, Degn K, Arango AS, Tan Z, Zhou J, et al. RosettaDDGPrediction "
  "for high-throughput mutational scans: from stability to binding. Protein Sci. "
  "2023;32(1):e4527. doi:10.1002/pro.4527", True),
 ("Steinegger M, Söding J. MMseqs2 enables sensitive protein sequence searching for the "
  "analysis of massive data sets. Nat Biotechnol. 2017;35(11):1026–1028. "
  "doi:10.1038/nbt.3988", True),
 ("Tsuboyama K, Dauparas J, Chen J, Laine E, Mohseni Behbahani Y, Weinstein JJ, et al. "
  "Mega-scale experimental analysis of protein folding stability in biology and design. "
  "Nature. 2023;620(7973):434–444. doi:10.1038/s41586-023-06328-6", True),
 ("Järvelin K, Kekäläinen J. Cumulated gain-based evaluation of IR techniques. ACM Trans "
  "Inf Syst. 2002;20(4):422–446. doi:10.1145/582415.582418", True),
 ("Efron B, Tibshirani RJ. An Introduction to the Bootstrap. New York: Chapman & Hall; "
  "1993.", True),
 ("Sokal RR, Rohlf FJ. Biometry: The Principles and Practice of Statistics in Biological "
  "Research. 3rd ed. New York: WH Freeman; 1995.", True),
 # 未能核实原始出处的方法论文 ⇒ 明确标注，不编造
 ("DDGPred: original method publication. REFERENCE_VERIFY_REQUIRED "
  "(per-mutation values were obtained from the benchmark table distributed with "
  "CATH-ddG [ref 4]; the primary DDGPred publication was not verified in this work.)",
  False),
 ("HERMES: original method publication. REFERENCE_VERIFY_REQUIRED "
  "(per-mutation values were obtained from the official HERMES experiment repository; "
  "the primary publication was a bioRxiv preprint and its final citation was not "
  "verified in this work.)", False),
]
n_unver = sum(1 for _, ok in REFS if not ok)
print(f"参考文献 {len(REFS)} 条，其中 REFERENCE_VERIFY_REQUIRED = {n_unver}")


# ---------------- helpers ----------------
def setup(doc, lang):
    st = doc.styles["Normal"]
    st.font.size = Pt(11 if not lang == "zh" else 10.5)
    st.font.name = "Times New Roman"
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.20
    pf.space_after = Pt(6)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.95)
        s.left_margin = s.right_margin = Inches(0.95)


def H(doc, text, lvl=1, keep=True):
    p = doc.add_heading(text, level=lvl)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)
        if lvl == 1:
            r.font.size = Pt(14.5)
        elif lvl == 2:
            r.font.size = Pt(11.5)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def P(doc, text, size=None, align="just", italic=False, bold=False, indent=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    p.alignment = (WD_ALIGN_PARAGRAPH.JUSTIFY if align == "just"
                   else WD_ALIGN_PARAGRAPH.LEFT if align == "left"
                   else WD_ALIGN_PARAGRAPH.CENTER)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    return p


def RICH(doc, parts, align="just", size=None):
    """parts: list of (text, style) where style in {'', 'i', 'b', 'bi', 'sub', 'sup'}"""
    p = doc.add_paragraph()
    for t, s in parts:
        r = p.add_run(t)
        r.italic = "i" in s
        r.bold = "b" in s
        if s == "sub":
            r.font.subscript = True
        if s == "sup":
            r.font.superscript = True
        if size:
            r.font.size = Pt(size)
    p.alignment = (WD_ALIGN_PARAGRAPH.JUSTIFY if align == "just"
                   else WD_ALIGN_PARAGRAPH.CENTER)
    return p


def EQ(doc, parts):
    """居中显示的公式行（真实上下标 + 斜体变量）。"""
    p = RICH(doc, parts, align="center")
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    return p


def cap(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(9.0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = False
    return p


def tbl(doc, df, caption_above=None, caption_below=None, fs=8.6):
    if caption_above:
        p = doc.add_paragraph()
        r = p.add_run(caption_above)
        r.font.size = Pt(9.0)
        r.bold = False
        p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    for i, c in enumerate(df.columns):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(c))
        run.bold = True
        run.font.size = Pt(fs)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(fs)
    if caption_below:
        cap(doc, caption_below)
    else:
        doc.add_paragraph()
    return t


# ---------------- main tables ----------------
T_overview = pd.DataFrame([
    {"Item": "Evaluation mutations (rows)", "Value": "4,473"},
    {"Item": "Protein targets", "Value": "214"},
    {"Item": "Structural families", "Value": "108"},
    {"Item": "Targets with n ≥ 5 mutations", "Value": "130"},
    {"Item": "Prediction configurations evaluated", "Value": "202"},
    {"Item": "Independent method families", "Value": "63"},
    {"Item": "Primary matched universe (rows)", "Value": f"{PU['n_rows']:,}"},
    {"Item": "Primary matched universe (targets n ≥ 5)", "Value": f"{PU['n_ge5']}"},
    {"Item": "Primary matched universe (method families)", "Value":
        f"{PU['n_method_families']}"},
    {"Item": "External methods with unambiguous mapping", "Value": "5"},
    {"Item": "External unambiguous mutations", "Value":
        f"{FR['external_clean']['n_rows']:,}"},
])
T_split = pd.DataFrame([
    {"Split regime": "Random", "Pooled Pearson": f"{H4['random']['overall_pearson']:.4f}",
     "C_B": f"{H4['random']['C_B']:+.4f}", "C_W": f"{H4['random']['C_W']:+.4f}",
     "Macro Spearman": f"{H4['random']['macro_spearman']:.4f}",
     "RMSE": f"{H4['random']['rmse']:.4f}"},
    {"Split regime": "Complex", "Pooled Pearson": f"{H4['complex']['overall_pearson']:.4f}",
     "C_B": f"{H4['complex']['C_B']:+.4f}", "C_W": f"{H4['complex']['C_W']:+.4f}",
     "Macro Spearman": f"{H4['complex']['macro_spearman']:.4f}",
     "RMSE": f"{H4['complex']['rmse']:.4f}"},
    {"Split regime": "seq40 (identical to complex)", "Pooled Pearson":
        f"{H4['seq40']['overall_pearson']:.4f}",
     "C_B": f"{H4['seq40']['C_B']:+.4f}", "C_W": f"{H4['seq40']['C_W']:+.4f}",
     "Macro Spearman": f"{H4['seq40']['macro_spearman']:.4f}",
     "RMSE": f"{H4['seq40']['rmse']:.4f}"},
    {"Split regime": "Family-OOD", "Pooled Pearson":
        f"{H4['family_OOD']['overall_pearson']:.4f}",
     "C_B": f"{H4['family_OOD']['C_B']:+.4f}", "C_W": f"{H4['family_OOD']['C_W']:+.4f}",
     "Macro Spearman": f"{H4['family_OOD']['macro_spearman']:.4f}",
     "RMSE": f"{H4['family_OOD']['rmse']:.4f}"},
])
T_ext = EP[["model", "n_rows", "n_targets_ge5", "overall_pearson", "macro_spearman",
            "C_B_row", "C_W", "rmse"]].copy()
T_ext["efficiency@10"] = EP["efficiency@10"]
T_ext.columns = ["Method", "n", "Targets (n≥5)", "Pooled Pearson", "Macro Spearman",
                 "C_B", "C_W", "RMSE", "Efficiency@10"]
for c in T_ext.columns[3:]:
    T_ext[c] = T_ext[c].round(4)


# ================================================================
def build(lang):
    zh = lang == "zh"
    T = (lambda en, zh_: zh_ if zh else en)
    SUF = "_ZH" if zh else "_EN"

    def FIG(doc, stem, width):
        doc.add_picture(str(FM / f"{stem}{SUF}.png"), width=Inches(width))

    D = Document()
    setup(D, lang)

    # ---------- Title page ----------
    t = D.add_paragraph()
    r = t.add_run(T("Disentangling Between-Target and Within-Target Signal in "
                    "PPI ΔΔG Benchmarking",
                    "解耦 PPI ΔΔG 基准评价中的跨靶信号与靶内信号"))
    r.bold = True
    r.font.size = Pt(19)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = D.add_paragraph()
    rs = sub.add_run(T("An exact covariance-signal decomposition under matched, "
                       "family-out-of-distribution evaluation",
                       "匹配、家族外推条件下的精确协方差信号分解"))
    rs.italic = True
    rs.font.size = Pt(11.5)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    D.add_paragraph()
    for line in [T("Author Name(s) [AUTHOR_PLACEHOLDER]", "作者姓名 [AUTHOR_PLACEHOLDER]"),
                 T("Affiliation(s) [AFFILIATION_PLACEHOLDER]",
                   "单位 [AFFILIATION_PLACEHOLDER]"),
                 T("Corresponding author: [NAME]  ·  Email: [EMAIL]  ·  "
                   "ORCID: [ORCID_PLACEHOLDER]",
                   "通讯作者：[NAME]  ·  邮箱：[EMAIL]  ·  ORCID：[ORCID_PLACEHOLDER]")]:
        p = D.add_paragraph()
        rr = p.add_run(line)
        rr.font.size = Pt(10.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    D.add_paragraph()
    p = D.add_paragraph()
    rr = p.add_run(T("Keywords: ", "关键词："))
    rr.bold = True
    rr.font.size = Pt(10)
    p.add_run(T("protein–protein interaction; binding affinity; ΔΔG; benchmarking; "
                "out-of-distribution generalisation; mutation effect prediction",
                "蛋白质–蛋白质相互作用；结合亲和力；ΔΔG；基准评价；"
                "分布外泛化；突变效应预测")).font.size = Pt(10)
    D.add_page_break()

    # ---------- Abstract (unstructured) ----------
    H(D, T("Abstract", "摘要"), 1)
    if zh:
        ab = (f"预测突变对蛋白质–蛋白质结合亲和力的影响，是抗体工程与界面设计的基础任务，"
              f"其进展高度依赖基准评价的可信度。然而常用的合并相关指标把两类性质不同的"
              f"预测信号混在一起：一类是把不同靶点放在正确结合量级上的跨靶信号，"
              f"另一类是在同一靶点内部把有利与不利突变正确排序的靶内信号。"
              f"我们引入一个精确的协方差信号分解，把合并 Pearson 相关写成 "
              f"r = C_B + C_W，其中 C_B 与 C_W 分别是跨靶与靶内协方差贡献；"
              f"并在严格匹配、来源筛选后的评价集合上应用该分解。"
              f"在仅改变划分方式的受控比较中（同一批 7,523 行突变、同一批模型配置），"
              f"合并 Pearson 中位数由随机划分的 {r0:.4f} 降至家族外推的 {r3:.4f}，"
              f"ΔC_B = {DL['delta_C_B']:+.4f}、ΔC_W = {DL['delta_C_W']:+.4f}；"
              f"按绝对变化计算，跨靶分量占下降的 {100*DL['fraction_between']:.1f}%，"
              f"而靶内分量仍为正。改变目标加权会改变对协方差贡献的解读："
              f"在五种加权方案下 C_B 仅在近零区间出现方向翻转（{sf0}/{len(W2)}），"
              f"而 C_W 在任何阈值下均不改变符号，但其幅度并非不变。"
              f"固定 K 短名单指标还受候选集几何约束：K=10 时 D(n,K)=K(n−K)/[n(n+1)] "
              f"在 n≤K 时为零，并在 n*≈{GEO['n_star_analytic']:.1f} 取峰，"
              f"本基准中 {100*GEO['n_le_K_fraction']:.1f}% 的 n≥5 靶点处于该退化区间。"
              f"在匹配的回溯消融中，该分解区分了仅改善合并分数的改动与同时改善靶内排序的改动。"
              f"外部可评价方法仅 5 个，功效不足，且其秩一致为正，因此不支持合并性能与"
              f"靶内性能的普遍解耦。我们建议 PPI ΔΔG 基准在明确声明的加权下"
              f"同时报告跨靶与靶内分量，并报告候选集几何。")
    else:
        ab = (f"Predicting how mutations alter protein–protein binding affinity underpins "
              f"antibody engineering and interface design, and progress depends on the "
              f"credibility of benchmark evaluation. Commonly reported pooled correlations, "
              f"however, conflate two qualitatively different sources of predictive signal: "
              f"a between-target component that places different targets on the correct "
              f"overall binding scale, and a within-target component that orders favourable "
              f"and unfavourable mutations correctly inside each target. We introduce an "
              f"exact covariance-signal decomposition that writes pooled Pearson "
              f"correlation as r = C_B + C_W, where C_B and C_W are the between-target and "
              f"within-target covariance contributions, and apply it to strictly matched, "
              f"provenance-filtered evaluation sets. In a controlled comparison that changes "
              f"only the split regime (identical 7,523 mutation rows and identical model "
              f"configurations), median pooled Pearson falls from {r0:.4f} under random "
              f"splitting to {r3:.4f} under family-out-of-distribution evaluation, with "
              f"ΔC_B = {DL['delta_C_B']:+.4f} and ΔC_W = {DL['delta_C_W']:+.4f}; on an "
              f"absolute-change basis the between-target component accounts for "
              f"{100*DL['fraction_between']:.1f}% of the decline while the within-target "
              f"component remains positive. Target weighting changes how these "
              f"contributions should be read: across five weighting schemes C_B reverses "
              f"direction in {sf0}/{len(W2)} configurations and does so only near zero, "
              f"whereas C_W keeps its sign at every threshold although its magnitude is not "
              f"invariant. Fixed-K shortlist metrics are further constrained by candidate-set "
              f"geometry: for K = 10 the discriminable interval "
              f"D(n,K) = K(n−K)/[n(n+1)] is zero when n ≤ K and peaks near "
              f"n* ≈ {GEO['n_star_analytic']:.1f}, and {100*GEO['n_le_K_fraction']:.1f}% of "
              f"targets with n ≥ 5 fall in that degenerate regime. In a matched retrospective "
              f"ablation the decomposition separated a change that improved pooled scores "
              f"only from one that also improved within-target ranking. External evidence "
              f"comprised five methods and is underpowered and positively aligned, so it does "
              f"not support a universal decoupling of pooled and within-target performance. "
              f"We recommend that PPI ΔΔG benchmarks report both components under an "
              f"explicitly declared weighting, together with candidate-set geometry.")
    P(D, ab)
    p = D.add_paragraph()
    rr = p.add_run(T("Keywords: ", "关键词："))
    rr.bold = True
    p.add_run(T("protein–protein interaction; binding affinity; ΔΔG; benchmarking; "
                "out-of-distribution generalisation; mutation effect prediction",
                "蛋白质–蛋白质相互作用；结合亲和力；ΔΔG；基准评价；"
                "分布外泛化；突变效应预测"))

    # ---------- Introduction ----------
    H(D, T("Introduction", "引言"), 1)
    if zh:
        P(D, "预测突变如何改变蛋白质–蛋白质结合亲和力（ΔΔG）是抗体亲和力成熟、"
             "界面优化与致病突变解释中的核心计算任务。随着深度学习方法的快速增加，"
             "该领域进展的判据越来越依赖基准评价：一个方法是否更好，通常由其在标准突变"
             "集合上报告的相关系数或误差决定。因此，评价协议本身的性质——它测量了什么、"
             "以及在何种外推条件下测量——直接决定了领域能得出什么结论。")
        P(D, "目前的评价实践以合并指标为主，即在整个突变集合上计算 Pearson 或 Spearman "
             "相关以及 RMSE；同时常辅以按复合物或按靶点计算的指标。近年已有工作揭示，"
             "这类评价对外推能力的刻画受到划分策略的强烈影响。Bushuiev 等量化了 PPI 数据中"
             "的结构性泄漏：按 PDB 代码划分时约 65% 的测试样本存在近重复训练样本，"
             "即使按 30% 序列同一性划分仍有约 30%，SKEMPI v2.0 若按 PPI 代码划分为 56%，"
             "而按领域专家分组则为 0%。CATH-ddG 进一步采用 CATH 同源超家族划分，"
             "并以最大 TM-score < 0.6 定义「难突变」。需要明确的是："
             "泄漏感知划分、家族外推评价、按靶点报告以及短名单式（Top-K）评价，"
             "均非本文提出。")
        P(D, "然而，一个关键问题仍未被直接回答：合并基准性能究竟来自哪里？"
             "把合并指标与按靶点指标并列报告，并不等于量化了跨靶信号与靶内信号对合并性能的"
             "加性贡献。一个模型可能主要因为能把不同靶点放在正确的整体量级上而获得较高的"
             "合并相关，而其在靶点内部的排序能力可能很弱；反之亦然。在靶点数量多、"
             "各靶点突变数高度不均的基准中，合并相关还会被大靶点主导，使这一混淆进一步加剧。")
        P(D, "本文引入一个精确的协方差信号分解来回答这一问题。我们把合并 Pearson 相关"
             "严格分解为跨靶分量 C_B 与靶内分量 C_W，使「合并性能由哪一部分构成」"
             "成为可直接测量的量，并在匹配、来源筛选后的评价集合上使用该分解量化三件事："
             "划分严格度改变的是哪一个分量；目标加权如何改变对这两个分量的解读；"
             "以及固定 K 短名单评价受候选集几何的何种约束。我们进一步在一个匹配的回溯消融中"
             "展示该分解能区分「仅改善合并分数」与「同时改善靶内排序」的模型改动，"
             "并在五个独立开发的外部方法上界定当前证据的泛化边界。")
    else:
        P(D, "Predicting how mutations alter protein–protein binding affinity (ΔΔG) is "
             "central to antibody affinity maturation, interface optimisation and the "
             "interpretation of disease-associated variants. As the number of deep-learning "
             "methods grows rapidly, progress in the field is increasingly judged by "
             "benchmark evaluation: whether one method is better is usually decided by the "
             "correlation or error it reports on a standard mutation set. The properties of "
             "the evaluation protocol itself — what it measures, and under which "
             "extrapolation conditions — therefore directly determine what the field can "
             "conclude.")
        P(D, "Current practice relies largely on pooled metrics, computing Pearson or "
             "Spearman correlation and RMSE over the entire mutation set, sometimes "
             "supplemented by per-complex or per-target metrics. Recent work has shown that "
             "how well such evaluation captures extrapolation depends strongly on the "
             "splitting strategy. Bushuiev et al. quantified structural leakage in PPI "
             "data: about 65% of test examples have a near-duplicate training example under "
             "PDB-code splitting, about 30% even under 30% sequence-identity splitting, 56% "
             "for SKEMPI v2.0 split by PPI code, and 0% under domain-expert grouping. "
             "CATH-ddG adopted CATH homologous-superfamily splits and defined 'hard' "
             "mutations by a maximum TM-score below 0.6. We emphasise that leakage-aware "
             "splitting, family-out-of-distribution evaluation, per-target reporting and "
             "shortlist (Top-K) evaluation were none of them introduced by this work.")
        P(D, "A key question nevertheless remains unanswered directly: where does pooled "
             "benchmark performance come from? Reporting pooled and per-target metrics side "
             "by side does not quantify the additive contribution of between-target and "
             "within-target signal to pooled performance. A model may achieve a high pooled "
             "correlation mainly because it places different targets on the correct overall "
             "scale, while its ability to order mutations within a target is weak — or the "
             "reverse. In benchmarks containing many targets with highly unequal mutation "
             "counts, pooled correlation is additionally dominated by large targets, which "
             "amplifies the conflation.")
        P(D, "Here we introduce an exact covariance-signal decomposition that addresses this "
             "question. We decompose pooled Pearson correlation rigorously into a "
             "between-target contribution C_B and a within-target contribution C_W, which "
             "makes 'what pooled performance is made of' a directly measurable quantity. "
             "Using this decomposition on matched, provenance-filtered evaluation sets, we "
             "quantify three things: which component is changed by split strictness; how "
             "target weighting changes the interpretation of the two components; and which "
             "constraints candidate-set geometry places on fixed-K shortlist evaluation. We "
             "further show, in a matched retrospective ablation, that the decomposition "
             "distinguishes model modifications that improve pooled scores only from those "
             "that also improve within-target ranking, and we bound the current evidence "
             "using five independently developed external methods.")

    # ---------- Results ----------
    H(D, T("Results", "结果"), 1)

    H(D, T("Pooled PPI ΔΔG performance contains separable between- and within-target signal",
           "合并 PPI ΔΔG 性能包含可分离的跨靶与靶内信号"), 2)
    if zh:
        P(D, f"在一个匹配的、通过预设来源标准的评价集合上"
             f"（{PU['n_rows']:,} 行突变、{PU['n_ge5']} 个 n≥5 靶点、"
             f"{PU['n_method_families']} 个独立方法族），我们把合并相关精确分解为两个"
             f"协方差贡献。对每个靶点 c，把真值与预测分别写成靶点均值与靶内偏差之和，"
             f"则合并相关满足以下恒等式，且残差在机器精度内为零（≤ 3.3×10⁻¹⁶）。"
             f"这一恒等式把一个通常被当作单一数值报告的指标，转换为两个可分别测量、"
             f"量纲一致的分量。需要强调的是，C_B 与 C_W 是协方差信号贡献，"
             f"不是独立相关系数、独立能力分数或因果机制。")
    else:
        P(D, f"On a matched evaluation set passing prespecified provenance criteria "
             f"({PU['n_rows']:,} mutation rows, {PU['n_ge5']} targets with n ≥ 5, "
             f"{PU['n_method_families']} independent method families), we decompose pooled "
             f"correlation exactly into two covariance contributions. Writing the true value "
             f"and the prediction for each mutation as a target mean plus a within-target "
             f"deviation, the pooled correlation satisfies the identity below, with residual "
             f"at machine precision (≤ 3.3 × 10⁻¹⁶). This identity converts a metric usually "
             f"reported as a single number into two separately measurable components on a "
             f"common scale. We stress that C_B and C_W are covariance-signal contributions, "
             f"not independent correlations, independent skill scores, or causal mechanisms.")
    EQ(D, [("y", "i"), ("ci", "sub"), (" = μ", "i"), ("c", "sub"), (" + ε", "i"),
           ("ci", "sub"), ("        ", ""), ("ŷ", "i"), ("ci", "sub"), (" = μ̂", "i"),
           ("c", "sub"), (" + ε̂", "i"), ("ci", "sub")])
    EQ(D, [("r", "i"), ("pooled", "sub"), (" = C", "i"), ("B", "sub"), (" + C", "i"),
           ("W", "sub")])
    FIG(D, "Fig1_capability_resolved_overview", 6.4)
    cap(D, ("图 1 | PPI ΔΔG 预测器的能力分辨评价。(a) 合并基准汇集多个靶点，"
            "每个靶点有其自身的平均 ΔΔG 水平与靶内突变变异。(b) 合并相关同时取决于"
            "靶点均值之间的差异（跨靶，C_B）与各靶点内部的突变排序（靶内，C_W）；"
            "面板内散点示意三个靶点。(c) 本研究所用评价流程。(d) 受控划分切换下的"
            "主结果：同一批 7,523 行与同一批配置，仅改变划分方式；"
            f"合并 r 由 {r0:.3f} 降至 {r3:.3f}，跨靶分量占绝对下降的 "
            f"{100*DL['fraction_between']:.1f}%。")
        if zh else
        ("Figure 1 | Capability-resolved evaluation of PPI ΔΔG predictors. "
         "(a) A pooled benchmark aggregates many targets, each with its own mean ΔΔG level "
         "and within-target mutation variability. (b) Pooled correlation depends both on "
         "differences between target means (between-target, C_B) and on the ordering of "
         "mutations within each target (within-target, C_W); points illustrate three "
         "stylised targets. (c) Evaluation workflow used in this study. (d) The main "
         "controlled result: identical 7,523 rows and identical configurations with only "
         f"the split regime changed; pooled r falls from {r0:.3f} to {r3:.3f}, and the "
         f"between-target component accounts for {100*DL['fraction_between']:.1f}% of the "
         "absolute decline."))

    H(D, T("Family-OOD evaluation disproportionately removes between-target signal",
           "家族外推评价不成比例地移除跨靶信号"), 2)
    if zh:
        P(D, f"为分离划分严格度的效应，我们在同一批 7,523 行突变与同一批模型配置上，"
             f"仅改变划分方式。结果见表 2 与图 2。合并 Pearson 中位数由随机划分的 "
             f"{r0:.4f} 降至家族外推的 {r3:.4f}。分解显示该下降并非均分："
             f"ΔC_B = {DL['delta_C_B']:+.4f}，ΔC_W = {DL['delta_C_W']:+.4f}。"
             f"按绝对变化口径 |ΔC_B|/(|ΔC_B|+|ΔC_W|) 计算，跨靶分量占下降的 "
             f"{100*DL['fraction_between']:.1f}%，而 C_W 在家族外推下仍为正"
             f"（{cw3:+.4f}）。因此，家族外推评价下合并相关的下降，主要由跨靶协方差贡献的"
             f"丧失承载，而靶内协方差保持为正。")
        P(D, "这一结果与「模型记住了靶点身份」不是同一命题：本研究没有进行身份记忆的"
             "干预实验，因此我们不作该因果陈述。更稳妥的表述是，该模式与「模型依赖"
             "不能跨家族迁移的靶点特异性信息」相一致。另外，复合物层面与 seq40 划分"
             "在本数据上给出完全相同的分配，因此我们不把它们作为彼此独立的比较。")
    else:
        P(D, "To isolate the effect of split strictness, we compared identical model "
             "configurations on the same 7,523 mutation rows while changing only the split "
             "regime (Table 2, Figure 2). Median pooled Pearson falls from {:.4f} under "
             "random splitting to {:.4f} under family-out-of-distribution evaluation. The "
             "decomposition shows that this decline is not uniform: ΔC_B = {:+.4f} and "
             "ΔC_W = {:+.4f}. On an absolute-change basis, "
             "|ΔC_B|/(|ΔC_B|+|ΔC_W|), the between-target component accounts for {:.1f}% of "
             "the decline, while C_W remains positive under family-out-of-distribution "
             "evaluation ({:+.4f}). Most of the reduction in pooled correlation under "
             "family-OOD evaluation is therefore carried by loss of the between-target "
             "covariance contribution, while within-target covariance remains positive."
             .format(r0, r3, DL["delta_C_B"], DL["delta_C_W"],
                     100 * DL["fraction_between"], cw3))
        P(D, "This is not the same proposition as 'the model memorised target identity': we "
             "performed no identity-memorisation intervention and therefore make no such "
             "causal claim. A more defensible reading is that the pattern is consistent with "
             "reliance on target-specific information that does not transfer across "
             "families. We also note that the complex-level and seq40 partitions yielded "
             "identical assignments in this dataset and were therefore not treated as "
             "independent comparisons.")
    FIG(D, "Fig2_controlled_familyOOD_decomposition", 6.5)
    cap(D, ("图 2 | 家族外推评价不成比例地移除合并相关中的跨靶分量。(a) 各划分方式下的"
            "合并 Pearson 中位数。(b) 精确分解为 C_B 与 C_W；C_W 在家族外推下仍为正。"
            f"(c) 绝对下降的构成：跨靶分量占 {100*DL['fraction_between']:.1f}%。"
            "四个划分方式使用同一批 7,523 行突变与同一批模型配置；复合物与 seq40 划分"
            "在本数据上相同，不构成独立比较。")
        if zh else
        ("Figure 2 | Family-out-of-distribution evaluation disproportionately removes the "
         "between-target component of pooled correlation. (a) Median pooled Pearson "
         "correlation by split regime. (b) Exact decomposition into C_B and C_W; C_W "
         "remains positive under family-OOD evaluation. (c) Composition of the absolute "
         f"decline, of which the between-target component accounts for "
         f"{100*DL['fraction_between']:.1f}%. All four regimes use the same 7,523 mutation "
         "rows and the same model configurations; the complex and seq40 partitions are "
         "identical in this dataset and are not independent comparisons."))
    tbl(D, T_split,
        caption_above=T("Table 2 | Controlled split comparison (identical rows and "
                        "configurations).",
                        "表 2 | 受控划分比较（同一批行、同一批配置）。"),
        caption_below=T("Median values across model configurations. The seq40 partition is "
                        "identical to the complex partition in this dataset.",
                        "跨模型配置的中位数。seq40 划分在本数据上与复合物划分完全相同。"))

    H(D, T("Target aggregation changes the interpretation of covariance contributions",
           "目标加权改变协方差贡献的解读"), 2)
    if zh:
        P(D, f"由于各靶点突变数高度不均，聚合权重的选择会影响两个分量的大小。"
             f"我们固定预测值，仅改变聚合权重，比较五种方案（每个突变等权、每个靶点等权、"
             f"权重正比于 √n、正比于 log(1+n)、以及上限为 30 的截断权重）。"
             f"结果显示，C_B 的方向不稳定性集中在近零区间：无阈值时 {sf0}/{len(W2)} 个配置"
             f"出现符号翻转，而在 |C_B| > 0.005 时不再出现（{sf5}/{len(W2)}）。"
             f"C_W 在任何阈值下均未出现符号翻转（{sfw}/{len(W2)}），但其幅度并非不变——"
             f"其极差中位数为 {W2.range_C_W.median():.4f}，大于 C_B 的 "
             f"{W2.range_C_B.median():.4f}。因此准确的表述是：C_B 仅在近零区间方向不稳，"
             f"而 C_W 保持符号稳定但幅度可变。任何报告的分解都必须声明其加权口径。")
    else:
        P(D, f"Because targets differ greatly in mutation count, the choice of aggregation "
             f"weight affects the magnitude of both components. Holding predictions fixed "
             f"and varying only the aggregation weights, we compared five schemes "
             f"(per-mutation, per-target, weight proportional to √n, proportional to "
             f"log(1+n), and capped at 30). C_B shows directional instability concentrated "
             f"near zero: {sf0}/{len(W2)} configurations reverse sign with no threshold, "
             f"whereas none do once |C_B| exceeds 0.005 ({sf5}/{len(W2)}). C_W never "
             f"reverses sign at any threshold ({sfw}/{len(W2)}), but its magnitude is not "
             f"invariant — its median range is {W2.range_C_W.median():.4f}, larger than the "
             f"{W2.range_C_B.median():.4f} of C_B. The accurate statement is therefore that "
             f"C_B is unstable in direction only near zero, whereas C_W keeps its sign but "
             f"varies in magnitude. Any reported decomposition must declare its weighting.")
    FIG(D, "Fig3_target_weighting_sensitivity", 6.2)
    cap(D, ("图 3 | 目标加权改变协方差贡献的解读。预测值不变，仅改变聚合权重。"
            "C_B 的方向不稳定性集中在近零区间；C_W 保持符号但幅度变化，"
            "故其幅度并非不变。")
        if zh else
        ("Figure 3 | Target weighting changes the interpretation of covariance "
         "contributions. Predictions are unchanged; only aggregation weights differ. C_B "
         "is unstable in direction only near zero, whereas C_W keeps its sign but varies "
         "in magnitude, so its magnitude is not invariant."))

    H(D, T("Fixed-K shortlist evaluation is constrained by candidate-set geometry",
           "固定 K 短名单评价受候选集几何约束"), 2)
    if zh:
        P(D, f"在突变文库或短名单筛选场景中，常用固定 K 的 Top-K 指标评价模型。"
             f"我们推导了这类指标的可区分区间。对候选集大小 n 与固定 K，"
             f"随机选取 K 条的期望收益占比为 K/n，而最优 K 条的收益占比为 "
             f"K(2n−K+1)/[n(n+1)]，因此可区分区间为 D(n,K) = K(n−K)/[n(n+1)]（n > K）；"
             f"当 n ≤ K 时 D = 0。也就是说，当靶点突变数不超过 K 时，任何模型都会选中"
             f"全部候选，该靶点对固定 K 短名单指标不携带任何信息。本基准中，"
             f"{GEO['n_ge5_count']} 个 n≥5 靶点里有 {GEO['n_le_K_count']} 个"
             f"（{100*GEO['n_le_K_fraction']:.1f}%）处于该退化区间。"
             f"此外，D 并非随 n 单调增加：它在 n* ≈ {GEO['n_star_analytic']:.1f} 取峰后"
             f"单调下降，例如 D(20) = {GEO['D_at_n20']:.4f} 而 D(50) = "
             f"{GEO['D_at_n50']:.4f}。因此「候选集越大越可区分」的启发式在本指标下不成立。")
        P(D, "上述几何结论不意味着 SKEMPI 的 ΔΔG 数值本身过于接近；"
             "后者需要以真实 ΔΔG 效用幅度做实证分析，本文未做该项分析。"
             "可辩护的表述是：当 n ≤ K 时固定 K 的候选选择评价在结构上不携带信息，"
             "且其可区分范围强烈依赖候选集几何。")
    else:
        P(D, f"Fixed-K top-K metrics are commonly used to evaluate models for mutation "
             f"library or shortlist screening. We derived the discriminable interval of "
             f"such metrics. For candidate-set size n and fixed K, the expected share of "
             f"gain from K randomly chosen candidates is K/n, whereas the optimal K "
             f"candidates attain K(2n−K+1)/[n(n+1)], so the discriminable interval is "
             f"D(n,K) = K(n−K)/[n(n+1)] for n > K, and D = 0 when n ≤ K. When a target has "
             f"no more than K mutations, any model selects every candidate and the target "
             f"carries no information for fixed-K shortlist metrics. In this benchmark, "
             f"{GEO['n_le_K_count']} of {GEO['n_ge5_count']} targets with n ≥ 5 "
             f"({100*GEO['n_le_K_fraction']:.1f}%) fall in that degenerate regime. "
             f"Moreover D is not monotonic in n: it peaks near "
             f"n* ≈ {GEO['n_star_analytic']:.1f} and declines thereafter, for example "
             f"D(20) = {GEO['D_at_n20']:.4f} versus D(50) = {GEO['D_at_n50']:.4f}. The "
             f"heuristic that larger candidate sets are more discriminable therefore does "
             f"not hold for this metric.")
        P(D, "This geometric result does not imply that SKEMPI ΔΔG values are intrinsically "
             "too similar; establishing that would require an analysis using actual ΔΔG "
             "utility magnitudes, which we did not perform. The defensible statement is "
             "that fixed-K candidate-selection evaluation is structurally non-informative "
             "when n ≤ K, and that its discriminable range depends strongly on "
             "candidate-set geometry.")
    FIG(D, "Fig4_fixed_K_candidate_geometry", 5.4)
    cap(D, ("图 4 | 固定 K 短名单评价受候选集几何约束。可区分区间 "
            "D(n,K)=K(n−K)/[n(n+1)] 在 n≤K 时为零，并在 n*≈20.5 取峰后下降；"
            f"本基准中 {100*GEO['n_le_K_fraction']:.1f}% 的 n≥5 靶点处于退化区间。")
        if zh else
        ("Figure 4 | Fixed-K shortlist evaluation is constrained by candidate-set "
         "geometry. The discriminable interval D(n,K) = K(n−K)/[n(n+1)] is zero for n ≤ K "
         "and peaks near n* ≈ 20.5 before declining; "
         f"{100*GEO['n_le_K_fraction']:.1f}% of targets with n ≥ 5 lie in the degenerate "
         "regime in this benchmark."))

    H(D, T("Capability decomposition clarifies retrospective remedy effects",
           "能力分解解释回溯性模型改动的效果"), 2)
    if zh:
        P(D, "最后，我们在一个匹配的消融集合中检验该分解的实用性。该集合中的模型共享"
             "同一 3,351 行宇宙，因此差异只来自模型改动本身。以「仅使用突变性质」的模型为"
             "基线：加入静态局部几何使合并 Pearson 提高约 0.028，但靶内 Macro Spearman "
             "几乎不变（−0.002），属于仅改善合并分数的改动；而引入「突变 × 配体 partner "
             "扰动」的改动同时改善两轴（合并 Pearson +0.070、Macro Spearman +0.044）；"
             "使用全量局部物理特征时两轴改善最大（+0.105 与 +0.076）。"
             "因此该分解能够区分「仅合并分数改善」与「同时改善靶内排序」两类改动。")
        P(D, "需要明确该分析的性质：它是回溯性的，即我们已知哪些改动有效，"
             "因此它说明的是该分解能够表达这一区别，而非基准能够前瞻预测哪条技术路线会成功。"
             "此外，本研究重新计算的 F4−F1 靶内差异（+0.044）与历史分析中的自助法区间"
             "（+0.073，[+0.021, +0.129]）基于不同的度量宇宙与聚合方式，"
             "两者不可在不加说明的情况下混用。")
    else:
        P(D, "Finally, we examined the practical utility of the decomposition in a matched "
             "ablation set whose models share an identical 3,351-row universe, so that "
             "differences arise only from the model modifications themselves. Relative to a "
             "baseline using mutation properties alone, adding static local geometry raised "
             "pooled Pearson by about 0.028 while leaving within-target Macro Spearman "
             "essentially unchanged (−0.002), a pooled-score-only improvement; introducing "
             "mutation × partner perturbation improved both axes (pooled Pearson +0.070, "
             "Macro Spearman +0.044); and a full local-physics model gave the largest "
             "improvement on both axes (+0.105 and +0.076). The decomposition therefore "
             "distinguishes a change that improves pooled scores only from one that also "
             "improves within-target ranking.")
        P(D, "The nature of this analysis should be stated plainly: it is retrospective, in "
             "that we already know which modifications were effective. It therefore shows "
             "that the decomposition can express this distinction, not that the benchmark "
             "prospectively predicts which technical route will succeed. In addition, our "
             "recomputed F4−F1 within-target difference (+0.044) and the bootstrap interval "
             "from the historical analysis (+0.073, [+0.021, +0.129]) rest on different "
             "metric universes and aggregation schemes and must not be mixed without "
             "explanation.")
    FIG(D, "Fig5_retrospective_remedy", 5.6)
    cap(D, ("图 5 | 匹配回溯消融中的能力分解。绿色表示两轴同时改善，红色表示仅合并分数改善"
            "或变化为负。该分析为回溯性，说明分解能够表达这一区别，"
            "不代表基准可前瞻预测成功路线。")
        if zh else
        ("Figure 5 | Capability decomposition in a matched retrospective ablation. Green "
         "indicates improvement on both axes; red indicates a pooled-score-only or negative "
         "change. The analysis is retrospective: it shows that the decomposition can express "
         "this distinction and does not imply prospective prediction of which modification "
         "will succeed."))

    H(D, T("External methods define the boundary of generalisation",
           "外部方法界定泛化边界"), 2)
    if zh:
        P(D, "为评估上述结论在独立开发的方法上的适用性，我们使用已公开的官方逐突变预测"
             "构建外部评价集合。在构建过程中我们发现，当匹配键省略链标识与野生型残基时，"
             "该键在标注侧并不唯一：4,473 行中仅 1,781 行具有唯一映射，另有 2,048 行为歧义键。"
             "因此主匹配评价只使用唯一映射行，歧义键被排除而非强行解析。"
             "这一处理使外部方法的表观合并相关系统性下降 0.02 至 0.12。")
        P(D, f"在唯一映射集合（{FR['external_clean']['n_rows']:,} 行、"
             f"{FR['external_clean']['n_ge5']} 个 n≥5 靶点）上，五个外部方法的总体相关与"
             f"靶内排序的排序一致，最优方法同为 {best.model}"
             f"（合并 Pearson {best.overall_pearson:+.4f}、Macro Spearman "
             f"{best.macro_spearman:+.4f}），即该集合内未出现方法层面的排名反转。"
             f"方法层面的秩一致为 ρ = "
             f"{er.get('rho_overall_pearson_vs_macro_spearman', float('nan')):+.2f}"
             f"（精确置换检验 p = "
             f"{er.get('perm_p_overall_pearson_vs_macro_spearman', float('nan')):.3f}）。"
             f"该结果为强正一致，因此不支持合并性能与靶内性能的普遍解耦；"
             f"同时它落在本研究内部观测到的异质范围内。仅有五个方法时，"
             f"方法层面的秩相关推断功效不足；按我们预设的实务阈值（8 个方法），"
             f"该分析被归类为功效不足。这是本研究的操作性选择，不是统计学通则。")
        P(D, "关于外部量的语义，我们特别说明：其中一个外部模型的官方输出为对数几率差"
             "（log_proba_mt − log_proba_wt），是无量纲的突变效应分数，"
             "而非已校准的 kcal/mol ΔΔG。因此我们不对该模型报告 RMSE 或 MAE，"
             "也不把其使用的变体标注为受结合亲和力监督训练的模型。")
    else:
        P(D, "To assess how far the above conclusions apply to independently developed "
             "methods, we built external evaluation sets from published official "
             "per-mutation predictions. During construction we found that when the matching "
             "key omits the chain identifier and the wild-type residue, the key is not "
             "unique on the annotation side: of 4,473 rows only 1,781 have a unique mapping "
             "and 2,048 are ambiguous. Primary matched evaluation therefore uses "
             "unambiguous rows only; ambiguous keys were excluded rather than forcibly "
             "resolved. This treatment systematically lowers apparent pooled correlation for "
             "external methods by 0.02 to 0.12.")
        P(D, f"On the unambiguous set ({FR['external_clean']['n_rows']:,} rows, "
             f"{FR['external_clean']['n_ge5']} targets with n ≥ 5), the five external "
             f"methods agree in their ordering on pooled correlation and on within-target "
             f"ranking, with the same best method, {best.model} "
             f"(pooled Pearson {best.overall_pearson:+.4f}, Macro Spearman "
             f"{best.macro_spearman:+.4f}); that is, no method-level rank reversal occurs "
             f"in this set. Method-level rank agreement is ρ = "
             f"{er.get('rho_overall_pearson_vs_macro_spearman', float('nan')):+.2f} (exact "
             f"permutation p = "
             f"{er.get('perm_p_overall_pearson_vs_macro_spearman', float('nan')):.3f}). "
             f"This is a strongly positive agreement, which does not support a universal "
             f"decoupling of pooled and within-target performance and falls within the "
             f"heterogeneous range observed internally. With only five methods, "
             f"method-level rank-correlation inference is low-powered; under our "
             f"prespecified practical threshold of eight methods the analysis is classified "
             f"as underpowered. This is a study-specific operating choice, not a general "
             f"statistical law.")
        P(D, "Regarding the semantics of one external quantity, we note that the official "
             "output of one model is a log-odds difference "
             "(log_proba_mt − log_proba_wt), a dimensionless mutation-effect score rather "
             "than a calibrated kcal/mol ΔΔG. We therefore report neither RMSE nor MAE for "
             "that model, and we do not label the variant we used as supervised on binding "
             "affinity.")
    tbl(D, T_ext,
        caption_above=T("Table 3 | External method capability profiles on the unambiguous "
                        "mapping set.",
                        "表 3 | 唯一映射集合上的外部方法能力画像。"),
        caption_below=T("All methods are evaluated on the same 1,781 unambiguous mutations. "
                        "RMSE is reported only for outputs with ΔΔG-like scaling.",
                        "所有方法在同一批 1,781 个唯一映射突变上评价。"
                        "仅对具有 ΔΔG 量纲的输出报告 RMSE。"))
    FIG(D, "Fig6_external_boundary", 5.4)
    cap(D, ("图 6 | 外部方法能力画像界定泛化边界。五个独立开发的方法在匹配且无歧义的映射上"
            "评价；秩一致为正，仅五个方法时功效不足，因此不支持普遍解耦。")
        if zh else
        ("Figure 6 | External method capability profiles define the boundary of "
         "generalisation. Five independently developed methods evaluated on a matched, "
         "unambiguous mapping; rank agreement is positive but the panel is underpowered "
         "with five methods, so a universal decoupling is not supported."))

    # ---------- Discussion ----------
    H(D, T("Discussion", "讨论"), 1)
    if zh:
        P(D, "本研究表明，在 PPI ΔΔG 基准中通常作为单一数值报告的合并相关，"
             "实际上由两个性质不同的预测信号构成：把不同靶点放在正确整体量级上的跨靶信号，"
             "以及在靶点内部正确排序突变效应的靶内信号。我们给出的分解是精确的恒等式，"
             "因此「合并性能由哪一部分构成」不再是解释性问题，而是可直接计算的量。"
             "这为评价结果的解读提供了此前缺失的一个维度。")
        P(D, "与既有工作的关系需要明确界定。PPI 数据中的结构性泄漏及其对划分策略的依赖"
             "已由 Bushuiev 等系统量化，他们同时提出了基于界面相似度的划分建议；"
             "CATH-ddG 采用 CATH 同源超家族划分，并以结构相似度定义了更难的外推突变；"
             "按复合物或按靶点报告指标以及短名单式评价在既有 ΔΔG 文献中亦已出现。"
             "因此本文的贡献不在于上述任何单一成分，而在于把精确的跨靶/靶内协方差分解"
             "与划分严格度、目标加权、匹配条件下的排名一致性以及候选集几何四项审计"
             "系统整合在同一套匹配、来源筛选的评价框架内，并辅以一个回溯性的模型改动归因案例。"
             "据我们所知，我们未发现既有工作在同等匹配条件下联合报告上述各项。")
        P(D, "关于 C_B 与 C_W 的解读，需要强调其边界。二者是同一恒等式中的协方差贡献，"
             "共享同一尺度，因此可以相加得到合并相关；它们不是两个独立估计的能力分数，"
             "也不应被解释为因果机制。我们特别避免把靶内分量的存在解读为「模型学会了物理」，"
             "或把跨靶分量的丧失解读为「模型记住了靶点标识」。"
             "一个可检验但本研究未做的方向是引入身份记忆或靶点打乱的干预实验，"
             "以直接判定跨靶分量的来源。")
        P(D, "加权敏感性审计给出一个对实践有直接影响的结论：跨靶分量的符号在近零区间"
             "对聚合权重敏感，而靶内分量的符号稳定。这意味着当报告的 C_B 接近零时，"
             "「跨靶信息是否存在」的结论可能取决于加权口径，而非取决于模型。"
             "因此我们建议：报告分解时必须同时声明加权方案，并在 C_B 接近零时"
             "同时给出多种加权下的结果。相反，跨靶分量绝对值较大的情况下该敏感性不构成实质风险。")
        P(D, "候选集几何的结果对 Top-K 评价实践有直接影响。当靶点突变数不超过 K 时，"
             "固定 K 的短名单指标在结构上无法区分模型，因为所有模型都会选中全部候选；"
             "本基准中约四成 n≥5 靶点处于该情形。同时可区分区间随候选集增大而先升后降，"
             "因此「扩大候选集以提高区分度」并非普遍有效。我们建议在报告短名单指标时"
             "同时给出 n 与 K、单独标注 n ≤ K 的靶点，或改用分数型（按比例 q）评价，"
             "以避免把结构性退化误读为模型能力差异。")
        P(D, "回溯性消融案例说明该框架在解释模型改动时的实用价值："
             "它能指出某项改动只提高了合并分数（可能主要来自跨靶分量的变化），"
             "而另一项改动同时改善了靶内排序。这一区分对方法设计具有直接意义，"
             "因为它把「分数提高」分解为可归因的两个来源。但该案例的回溯性质必须明确："
             "它不构成对模型改进路线的前瞻预测能力。")
        P(D, "最后，外部证据界定了当前结论的泛化边界。在五个独立开发的方法上，"
             "合并性能与靶内排序的方法排序一致（强正一致），这既不支持也不反驳"
             "「两轴普遍解耦」的主张；同时该面板功效不足，样本量远小于稳健推断所需。"
             "因此我们把外部结果表述为「与异质能力画像相容，但对普遍解耦功效不足」，"
             "而不作方向性结论。未来工作应扩大外部方法面板，并在独立的前瞻实验数据集上"
             "检验分解分量与实际筛选收益之间的关系。")
    else:
        P(D, "This study shows that the pooled correlation conventionally reported as a "
             "single number in PPI ΔΔG benchmarking is in fact composed of two qualitatively "
             "different predictive signals: a between-target component that places different "
             "targets on the correct overall scale, and a within-target component that orders "
             "mutation effects correctly inside each target. Because our decomposition is an "
             "exact identity, 'what pooled performance is made of' ceases to be a matter of "
             "interpretation and becomes directly computable. This adds a previously missing "
             "dimension to the reading of benchmark results.")
        P(D, "The relation to prior work should be delimited clearly. Structural leakage in "
             "PPI data and its dependence on splitting strategy were quantified "
             "systematically by Bushuiev et al., who also proposed interface-similarity "
             "splitting; CATH-ddG adopted CATH homologous-superfamily splits and used "
             "structural similarity to define harder extrapolation mutations; and "
             "per-complex or per-target reporting together with shortlist-style evaluation "
             "already appear in the ΔΔG literature. Our contribution therefore lies not in "
             "any single one of these ingredients but in the systematic integration of an "
             "exact between/within-target covariance decomposition with four audits — split "
             "strictness, target weighting, matched rank agreement, and candidate-set "
             "geometry — inside one matched, provenance-filtered framework, together with a "
             "retrospective attribution case study. To our knowledge we did not identify "
             "prior work that jointly reports these under comparable matching conditions.")
        P(D, "The interpretation of C_B and C_W has clear boundaries. They are covariance "
             "contributions within one identity and share a common scale, so they sum to "
             "pooled correlation; they are not two independently estimated skill scores and "
             "should not be read as causal mechanisms. In particular we avoid interpreting "
             "the presence of within-target signal as evidence that a model has learned "
             "physics, or the loss of between-target signal as evidence that a model "
             "memorised target identity. A directly testable but unexplored direction is an "
             "intervention that perturbs target identity or shuffles targets, which would "
             "establish the origin of the between-target component.")
        P(D, "The weighting-sensitivity audit yields a conclusion of direct practical "
             "relevance: the sign of the between-target component is sensitive to "
             "aggregation weights near zero, whereas the sign of the within-target "
             "component is stable. Consequently, when a reported C_B is close to zero, the "
             "conclusion about whether between-target information is present may depend on "
             "the weighting convention rather than on the model. We therefore recommend "
             "that any reported decomposition declare its weighting scheme, and that "
             "multiple weightings be shown when C_B is near zero. Where the absolute value "
             "of C_B is large, this sensitivity does not constitute a material risk.")
        P(D, "The candidate-set geometry result has direct implications for Top-K "
             "evaluation practice. When a target has no more than K mutations, fixed-K "
             "shortlist metrics cannot distinguish models structurally, because every model "
             "selects all candidates; roughly four in ten targets with n ≥ 5 fall in this "
             "regime in our benchmark. Because the discriminable interval first rises and "
             "then falls as candidate sets grow, enlarging a candidate set is not "
             "universally effective at improving discrimination. We recommend reporting n "
             "alongside K, labelling n ≤ K targets separately, or adopting fraction-based "
             "(q) evaluation, so that structural degeneracy is not misread as a difference "
             "in model capability.")
        P(D, "The retrospective ablation illustrates the framework's explanatory value for "
             "model modifications: it identifies a change that raised pooled scores only "
             "— plausibly through the between-target component — as distinct from a change "
             "that also improved within-target ranking. This distinction matters for method "
             "design because it attributes an improvement to one of two sources. The "
             "retrospective nature of the case study must, however, be stated explicitly: it "
             "does not establish prospective ability to predict which modifications will work.")
        P(D, "Finally, external evidence bounds the generalisation of these conclusions. "
             "Across five independently developed methods the ordering of methods agreed "
             "between pooled and within-target performance, a strongly positive agreement "
             "that neither supports nor refutes a universal decoupling of the two axes; the "
             "panel is also underpowered, with far fewer methods than robust inference would "
             "require. We therefore report the external result as compatible with "
             "heterogeneous capability profiles but underpowered with respect to universal "
             "decoupling, and draw no directional conclusion from it. Future work should "
             "broaden the external panel and test the relationship between the decomposed "
             "components and realised screening yield on independent prospective datasets.")

    # ---------- Limitations ----------
    H(D, T("Limitations", "局限"), 1)
    if zh:
        P(D, "本研究的外推设定基于靶点家族划分，但这并不能排除基础模型在预训练阶段"
             "接触过与评价家族相似的序列或结构；我们只能控制任务特异性的监督训练重叠，"
             "无法排除预训练层面的相似性。")
        P(D, "基准中包含的模型池不是全部 PPI ΔΔG 方法的随机样本，且部分模型族的覆盖有限；"
             "外部可评价方法仅五个，方法层面的推断功效不足。因此本文关于能力构成的结论"
             "应被理解为一组受控观测，而非对领域整体的普查。")
        P(D, "短名单与候选选择分析均以已有预测数据的重新分析为基础，不包含前瞻性湿实验验证；"
             "在靶点突变数不超过 K 时，相关指标本身不携带信息，因此与选择相关的结论"
             "在本文中仅作低置信度描述。回溯性消融案例同样不构成对改进路线的前瞻预测。")
        P(D, "C_B 与 C_W 是协方差信号贡献，不是因果机制；本文未进行身份记忆或靶点打乱等"
             "干预实验，因此不对跨靶分量的来源作因果断言。")
        P(D, "本研究没有独立的、前瞻性的实验数据集，因此未进行外部数据集验证；"
             "同时本文不提供通用模型推荐器，也不主张任何方法族具有普遍优势。")
    else:
        P(D, "Our extrapolation setting is based on target-family partitioning, which does "
             "not exclude the possibility that foundation-model pretraining encountered "
             "sequences or structures similar to the evaluation families. We can control "
             "task-specific supervised overlap but not pretraining-level similarity.")
        P(D, "The model pool assembled here is not a random sample of all PPI ΔΔG methods, "
             "and coverage of some model families is limited; only five methods are "
             "externally evaluable, so method-level inference is underpowered. Conclusions "
             "about the composition of capability should therefore be read as controlled "
             "observations rather than a survey of the field.")
        P(D, "Shortlist and candidate-selection analyses rest on re-analysis of existing "
             "predictions and include no prospective wet-lab validation; when a target has "
             "no more than K mutations the relevant metrics carry no information, so "
             "selection-related conclusions are reported here as low-confidence "
             "descriptions only. The retrospective ablation likewise does not establish "
             "prospective prediction of which modifications will succeed.")
        P(D, "C_B and C_W are covariance-signal contributions rather than causal mechanisms. "
             "We performed no identity-memorisation or target-shuffling intervention and "
             "therefore make no causal claim about the origin of the between-target component.")
        P(D, "No independent prospective experimental dataset was available, so no external "
             "dataset validation was performed. This work does not provide a universal model "
             "recommender and does not claim general superiority of any method family.")

    # ---------- Data / contributions / ack ----------
    H(D, T("Data and code availability", "数据与代码可用性"), 1)
    P(D, T("All analyses use existing predictions and published external predictions. "
           "Frozen input hashes, the machine-readable source of every headline number, and "
           "an automated consistency checker are provided with the Supplementary "
           "Information so that all reported values can be regenerated without retraining "
           "any model. No new model was trained and no GPU was used.",
           "本分析仅使用既有预测与已公开的外部预测。冻结输入哈希、所有头条数字的机器可读"
           "来源以及自动化一致性检查器随补充材料提供，使全部报告数值可在不重新训练任何模型的"
           "情况下重新生成。本研究未训练新模型，也未使用 GPU。"))

    H(D, T("Author contributions", "作者贡献"), 1)
    P(D, T("Study conception, analysis design, interpretation of results and writing of "
           "the manuscript were performed by the author(s). The use of AI-assisted tools is "
           "declared below.",
           "研究构思、分析设计、结果解读与稿件撰写由作者完成；AI 辅助工具的使用情况"
           "见下节声明。"))

    H(D, T("Acknowledgements", "致谢"), 1)
    P(D, T("[FUNDING_PLACEHOLDER] We thank the developers of the publicly released "
           "prediction sets and datasets used in this study.",
           "[FUNDING_PLACEHOLDER] 我们感谢本研究所用公开预测集合与数据集的开发者。"))

    H(D, T("Use of AI-assisted tools", "AI 辅助工具使用声明"), 1)
    P(D, T("AI-assisted coding and language tools were used to support code drafting and "
           "debugging, generation of plotting scripts, document automation, "
           "literature/resource retrieval, and language editing. All metric definitions, "
           "analysis choices, interpretation of results, scientific claims, and final "
           "manuscript content were reviewed and approved by the authors, who take full "
           "responsibility for the work. AI tools were not used to generate experimental "
           "data or to modify upstream model predictions.",
           "本研究使用 AI 辅助编码与语言工具，用于支持代码草拟与调试、绘图脚本生成、"
           "文档自动化、文献与资源检索以及语言润色。所有指标定义、分析选择、结果解读、"
           "科学主张与最终稿件内容均由作者审阅并确认，作者对全部内容负责。"
           "AI 工具未被用于生成实验数据或修改任何上游模型预测。"))
    P(D, T("A more detailed technical disclosure, including the exact scope of automated "
           "generation, is provided in the Supplementary Information.",
           "更详细的技术披露（包括自动化生成的具体范围）见补充材料。"))

    # ---------- References ----------
    H(D, T("References", "参考文献"), 1)
    for i, (ref, ok) in enumerate(REFS, 1):
        p = D.add_paragraph()
        run = p.add_run(f"{i}. {ref}")
        run.font.size = Pt(9.5)
        if not ok:
            run.font.color.rgb = RGBColor(0x99, 0x00, 0x00)
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(3)

    D.add_page_break()
    H(D, T("Table 1 | Dataset and evaluation overview", "表 1 | 数据集与评价概览"), 1)
    tbl(D, T_overview, fs=9.2)

    name = (f"CR-PPI-ddG_manuscript_v2_{'ZH' if zh else 'EN'}.docx")
    path = OUT / name
    D.save(str(path))
    print(f"  {name}  ({path.stat().st_size/1024:.0f} KB)")
    return path


zh_p = build("zh")
en_p = build("en")
print("v2 DOCX 完成")
